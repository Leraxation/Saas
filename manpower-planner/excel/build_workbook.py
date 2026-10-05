from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter as CL
import sys

OUT = sys.argv[1]
import os
NROWS = int(os.environ.get('NROWS', 5000))            # employee rows on the Data sheet
DEPT0, DEPTN = 5, 5 + int(os.environ.get('NDEPT', 50)) - 1    # 50 department rows
PLAN0, PLANN = 5, 5 + int(os.environ.get('NPLAN', 300)) - 1   # 300 plan rows
LAST = NROWS + 1

F = "Arial"
def font(**k): return Font(name=F, **k)
BLUE = "0000FF"; GREEN = "008000"
fill = lambda c: PatternFill("solid", fgColor=c)
YEL, INHDR, CALCHDR, CALC, TITLE, SUB, OKF, BADF, WARNF = fill("FFF59D"), fill("DCE6F2"), fill("E7E6E6"), fill("F5F5F5"), fill("1F3864"), fill("D9E1F2"), fill("E2EFDA"), fill("F8CBAD"), fill("FFE699")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
OMR = '#,##0;(#,##0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
HIDE0 = '#,##0;(#,##0);;'
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

def title(ws, text, sub, width_cols=8):
    ws["A1"] = text; ws["A1"].font = font(bold=True, size=16, color="FFFFFF"); ws["A1"].fill = TITLE
    for c in range(2, width_cols + 1): ws.cell(1, c).fill = TITLE
    ws["A2"] = sub; ws["A2"].font = font(italic=True, size=10, color="595959")
    ws.row_dimensions[1].height = 26
    ws.sheet_view.showGridLines = False

def hdr(cell, text, kind="calc"):
    cell.value = text
    cell.font = font(bold=True, size=10, color="000000")
    cell.fill = INHDR if kind == "input" else CALCHDR
    cell.alignment = CENTER; cell.border = BOX

# ------------------------------------------------------------------ Guide
g = wb.active; g.title = "Guide"
title(g, "Manpower Budget Workbook", "Every figure is calculated from the employee data you paste into the Data sheet. This template contains no data.", 6)
rows = [
 ("How to use it", None),
 ("1", "Settings: set the budget year and check every assumption. Mark each one Confirmed once Payroll or Finance has agreed it."),
 ("2", "Departments: type each department exactly as it is written in your HR export, with its overtime %, last year's budget and ceiling."),
 ("3", "Data: paste your employee export as values into columns A to P, starting at row 2 (Home > Paste > Paste Values). Do not type in the grey columns: they calculate themselves."),
 ("4", "Data, column AL (Check): fix every row it flags. Overview shows how many rows still have issues."),
 ("5", "Plan: add planned new hires and other exits that are not in the employee data."),
 ("6", "Overview: the budget, budget bridge, reconciliation checks and charts update automatically. Budget by Dept and Monthly give the detail."),
 ("", None),
 ("Colour legend", None),
 ("Blue text", "Something you type or paste (inputs)."),
 ("Yellow cell", "A key assumption that must be confirmed before the budget is submitted."),
 ("Grey column", "Calculated. Do not type here."),
 ("Green text", "Pulled from another sheet."),
 ("", None),
 ("Rules the workbook follows", None),
 ("Data first", "Nothing is filled in for you. Headcount, pay, departments and leavers all come from the Data sheet."),
 ("Traceable", "Every department figure is a SUMIFS or COUNTIFS over the Data sheet. Filter the Data sheet by department to see the rows behind any number."),
 ("Reconciled", "Overview compares the totals by department with the raw Data totals. Any difference (for example a department missing from the Departments sheet) is shown in red."),
 ("Capacity", f"Up to {NROWS:,} employees, {DEPTN-DEPT0+1} departments and {PLANN-PLAN0+1} plan lines."),
 ("", None),
 ("Data sheet format: one example row (do not paste this into the Data sheet)", None),
]
r = 4
for a, b in rows:
    if b is None:
        g.cell(r, 1, a).font = font(bold=True, size=12, color="1F3864")
    else:
        g.cell(r, 1, a).font = font(bold=True, size=10)
        c = g.cell(r, 2, b); c.font = font(size=10); c.alignment = WRAP
        if a == "Blue text": g.cell(r, 1).font = font(bold=True, size=10, color=BLUE)
        if a == "Yellow cell": g.cell(r, 1).fill = YEL
        if a == "Grey column": g.cell(r, 1).fill = CALC
        if a == "Green text": g.cell(r, 1).font = font(bold=True, size=10, color=GREEN)
    r += 1
DATA_COLS = ["Employee ID", "Employee name", "Department *", "Section", "Job title", "Grade", "Nationality *", "Gender", "Date of birth", "Joining date", "Contract end date", "Basic salary (OMR / month) *", "Allowances total (OMR / month)", "Employment status", "Planned exit month (1-12, optional)", "Personal increment % (optional)"]
example = ["E10234", "Example Employee", "Engineering & Maintenance", "Line Maintenance", "Licensed Engineer B1", "T3", "Omani", "Male", "14/03/1985", "01/09/2012", "", 1350, 650, "Active", "", ""]
for i, (h, v) in enumerate(zip(DATA_COLS, example)):
    hdr(g.cell(r, 1 + i), h, "input"); c = g.cell(r + 1, 1 + i, v); c.font = font(size=10, color=BLUE); c.border = BOX
g.cell(r + 2, 1, "Dates must be real Excel dates (not text). Nationality: Omani, Oman, OM or Local counts as Omani; anything else is expatriate. Status containing resign, terminat, inactive, vacan, deceased or separated is treated as not active.").font = font(size=9, italic=True)
g.column_dimensions["A"].width = 22; g.column_dimensions["B"].width = 95
for i in range(3, 17): g.column_dimensions[CL(i)].width = 16

# ------------------------------------------------------------------ Settings
s = wb.create_sheet("Settings")
title(s, "Settings and assumptions", "Assumptions are not data. Each one is shown here, used by name in the formulas, and must be confirmed.", 5)
for c, h in zip("ABCDE", ["Assumption", "Value", "Unit", "Confirmed?", "Source / note"]): hdr(s[f"{c}4"], h, "input")
SET = [
 ("FY", "Budget year", 2027, "year", "Fiscal year being budgeted."),
 ("IncPct", "Annual increment", 0.03, "% of basic and allowances", "Default 3%. Confirm with Compensation."),
 ("IncMonth", "Increment effective month", 1, "month 1-12", "1 = January."),
 ("BonusMonths", "Bonus / incentive", 1, "months of basic", "Accrued evenly."),
 ("SI_O", "Employer social protection, Omani", 0.125, "% of contribution base", "Default 12.5%. Confirm the current Social Protection Fund schedule with Payroll."),
 ("SI_E", "Employer contribution, expatriate", 0.01, "% of contribution base", "Default 1% (occupational injury). Confirm with Payroll."),
 ("SI_Base", "Contribution base", "Basic + allowances", "list", "Basic + allowances, or Basic only."),
 ("EOSB", "Expatriate end-of-service accrual", 0.0833, "% of basic", "8.33% = one month's basic per year. Confirm against the Labour Law and contracts."),
 ("Ticket", "Expatriate annual air ticket", 350, "OMR per head per year", "Confirm with Travel / HR policy."),
 ("MedO", "Medical insurance, Omani", 280, "OMR per head per year", "Confirm with the insurer's current premium."),
 ("MedE", "Medical insurance, expatriate", 420, "OMR per head per year", "Confirm with the insurer's current premium."),
 ("Training", "Training", 220, "OMR per head per year", "Excludes type ratings and licences budgeted elsewhere."),
 ("RecO", "Recruitment cost, Omani hire", 400, "OMR per hire", "Advertising, assessment, medicals."),
 ("RecE", "Recruitment cost, expatriate hire", 1800, "OMR per hire", "Agency, visa, medical, relocation."),
 ("Contingency", "Contingency", 0.02, "% of department total", ""),
 ("OmanTarget", "Omanisation target", 0.85, "% of closing headcount", "Company target."),
 ("RetAgeO", "Retirement age, Omani", 60, "years", "Used to find retirements in the budget year."),
 ("RetAgeE", "Retirement age, expatriate", 60, "years", ""),
 ("RetireExit", "Treat retirements as exits", "No", "Yes / No", "Yes = people reaching retirement age are paid only until their birth month."),
 ("ContractExit", "Treat contract ends as exits", "No", "Yes / No", "Yes = people whose contract ends in the year are paid only until that month."),
]
dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False); s.add_data_validation(dv_yes)
dv_base = DataValidation(type="list", formula1='"Basic + allowances,Basic only"'); s.add_data_validation(dv_base)
for i, (name, label, val, unit, note) in enumerate(SET):
    rr = 5 + i
    s.cell(rr, 1, label).font = font(size=10)
    v = s.cell(rr, 2, val); v.font = font(size=10, color=BLUE); v.fill = YEL; v.border = BOX
    if isinstance(val, float) and val < 1: v.number_format = "0.00%"
    s.cell(rr, 3, unit).font = font(size=9, color="595959")
    cf = s.cell(rr, 4, "No"); cf.font = font(size=10, color=BLUE); cf.border = BOX; dv_yes.add(cf)
    s.cell(rr, 5, note).font = font(size=9, color="595959"); s.cell(rr, 5).alignment = WRAP
    if name in ("RetireExit", "ContractExit"): dv_yes.add(v)
    if name == "SI_Base": dv_base.add(v)
    wb.defined_names[name] = DefinedName(name, attr_text=f"Settings!$B${rr}")
nset = len(SET)
SI_G_ROW = 5 + nset + 1
s.cell(SI_G_ROW, 1, "Contribution base includes allowances (1 = yes)").font = font(size=9, italic=True, color="595959")
s.cell(SI_G_ROW, 2, '=IF(SI_Base="Basic only",0,1)').font = font(size=9)
wb.defined_names["SI_Gross"] = DefinedName("SI_Gross", attr_text=f"Settings!$B${SI_G_ROW}")
s.cell(SI_G_ROW + 1, 1, "Assumptions not yet confirmed").font = font(bold=True, size=10)
s.cell(SI_G_ROW + 1, 2, f'=COUNTIF(D5:D{4+nset},"<>Yes")').font = font(bold=True, size=10)
wb.defined_names["Unconfirmed"] = DefinedName("Unconfirmed", attr_text=f"Settings!$B${SI_G_ROW+1}")
s.conditional_formatting.add(f"D5:D{4+nset}", CellIsRule(operator="equal", formula=['"Yes"'], fill=OKF))
s.conditional_formatting.add(f"D5:D{4+nset}", CellIsRule(operator="notEqual", formula=['"Yes"'], fill=WARNF))
for c, w in zip("ABCDE", [40, 20, 24, 13, 70]): s.column_dimensions[c].width = w
s.freeze_panes = "A5"

# ------------------------------------------------------------------ Departments (inputs + data-driven results)
d = wb.create_sheet("Departments")
title(d, "Departments", "Type the departments exactly as they appear in the Data sheet. Everything from column E onwards is calculated from the data.", 22)
DH = [("Department", "input"), ("Overtime % of basic", "input"), ("Prior-year budget (OMR)", "input"), ("FY ceiling (OMR)", "input"),
      ("Headcount today", ""), ("Omani today", ""), ("Omani % today", ""), ("Women", ""), ("Monthly gross today (OMR)", ""),
      ("Named leavers in FY", ""), ("Reaching retirement age", ""), ("Contracts ending", ""), ("Planned hires", ""), ("Planned other exits", ""),
      ("Closing headcount", ""), ("Staff FY cost (OMR)", ""), ("Plan lines cost (OMR)", ""), ("Contingency (OMR)", ""), ("FY budget (OMR)", ""),
      ("Change vs prior year (OMR)", ""), ("Ceiling status", ""), ("Closing Omani %", "")]
for i, (h, k) in enumerate(DH): hdr(d.cell(4, 1 + i), h, k)
D = lambda col: f"Data!${col}$2:${col}${LAST}"
P = lambda col: f"Plan!${col}${PLAN0}:${col}${PLANN}"
for rr in range(DEPT0, DEPTN + 1):
    for c in range(1, 5):
        cell = d.cell(rr, c); cell.font = font(size=10, color=BLUE); cell.border = BOX
    d.cell(rr, 2).number_format = "0.0%"; d.cell(rr, 3).number_format = OMR; d.cell(rr, 4).number_format = OMR
    A = f"$A{rr}"
    f = {
     5: f'=IF({A}="","",COUNTIFS({D("C")},{A},{D("R")},1))',
     6: f'=IF({A}="","",COUNTIFS({D("C")},{A},{D("R")},1,{D("Q")},1))',
     7: f'=IF(OR({A}="",N(E{rr})=0),"",F{rr}/E{rr})',
     8: f'=IF({A}="","",COUNTIFS({D("C")},{A},{D("R")},1,{D("H")},"F*"))',
     9: f'=IF({A}="","",SUMIFS({D("AK")},{D("C")},{A}))',
     10: f'=IF({A}="","",COUNTIFS({D("C")},{A},{D("R")},1,{D("W")},"<12"))',
     11: f'=IF({A}="","",COUNTIFS({D("C")},{A},{D("U")},">0"))',
     12: f'=IF({A}="","",COUNTIFS({D("C")},{A},{D("V")},">0"))',
     13: f'=IF({A}="","",SUMIFS({P("E")},{P("B")},{A},{P("A")},"Hire"))',
     14: f'=IF({A}="","",SUMIFS({P("E")},{P("B")},{A},{P("A")},"Exit"))',
     15: f'=IF({A}="","",E{rr}-J{rr}+M{rr}-N{rr})',
     16: f'=IF({A}="","",SUMIFS({D("AH")},{D("C")},{A}))',
     17: f'=IF({A}="","",SUMIFS({P("S")},{P("B")},{A}))',
     18: f'=IF({A}="","",(P{rr}+Q{rr})*Contingency)',
     19: f'=IF({A}="","",P{rr}+Q{rr}+R{rr})',
     20: f'=IF(OR({A}="",N(C{rr})=0),"",S{rr}-C{rr})',
     21: f'=IF(OR({A}="",N(D{rr})=0),"",IF(S{rr}>D{rr},"Over by "&TEXT(S{rr}-D{rr},"#,##0"),"Within ceiling"))',
     22: f'=IF(OR({A}="",N(O{rr})=0),"",(F{rr}-COUNTIFS({D("C")},{A},{D("R")},1,{D("Q")},1,{D("W")},"<12")+SUMIFS({P("E")},{P("B")},{A},{P("A")},"Hire",{P("D")},"Omani")-SUMIFS({P("E")},{P("B")},{A},{P("A")},"Exit",{P("D")},"Omani"))/O{rr})',
    }
    for c, fx in f.items():
        cell = d.cell(rr, c, fx); cell.font = font(size=10); cell.fill = CALC; cell.border = BOX
        cell.number_format = PCT if c in (7, 22) else (OMR if c in (9, 16, 17, 18, 19, 20) else INT)
TR = DEPTN + 1
d.cell(TR, 1, "Total").font = font(bold=True, size=10)
for c in [3, 4] + list(range(5, 7)) + list(range(8, 20)) + [20]:
    L = CL(c); cell = d.cell(TR, c, f"=SUM({L}{DEPT0}:{L}{DEPTN})"); cell.font = font(bold=True, size=10); cell.border = BOX
    cell.number_format = OMR if c in (3, 4, 9, 16, 17, 18, 19, 20) else INT
d.cell(TR, 7, f"=IF(E{TR}=0,\"\",F{TR}/E{TR})").number_format = PCT
d.cell(TR, 22, f'=IF(O{TR}=0,"",(F{TR}-COUNTIFS({D("R")},1,{D("Q")},1,{D("W")},"<12")+SUMIFS({P("E")},{P("A")},"Hire",{P("D")},"Omani")-SUMIFS({P("E")},{P("A")},"Exit",{P("D")},"Omani"))/O{TR})').number_format = PCT
for c in (7, 22): d.cell(TR, c).font = font(bold=True, size=10)
d.conditional_formatting.add(f"U{DEPT0}:U{DEPTN}", FormulaRule(formula=[f'LEFT(U{DEPT0},4)="Over"'], fill=BADF))
d.conditional_formatting.add(f"G{DEPT0}:G{DEPTN}", FormulaRule(formula=[f'AND(G{DEPT0}<>"",G{DEPT0}<OmanTarget)'], fill=WARNF))
d.conditional_formatting.add(f"V{DEPT0}:V{DEPTN}", FormulaRule(formula=[f'AND(V{DEPT0}<>"",V{DEPT0}<OmanTarget)'], fill=WARNF))
d.column_dimensions["A"].width = 32
for c in range(2, 23): d.column_dimensions[CL(c)].width = 14
d.row_dimensions[4].height = 44
d.freeze_panes = "B5"
dv_ot = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", showErrorMessage=True, error="Enter a percentage between 0% and 100%.")
d.add_data_validation(dv_ot); dv_ot.add(f"B{DEPT0}:B{DEPTN}")
d.cell(3, 1, "Departments found in Data but missing here:").font = font(size=9, italic=True)
d.cell(3, 5, f'=COUNTIF({D("AL")},"Department not on Departments sheet")').font = font(bold=True, size=10, color="C00000")

# ------------------------------------------------------------------ Data
x = wb.create_sheet("Data")
CALC_COLS = [
 ("Q", "Omani (1/0)"), ("R", "Active (1/0)"), ("S", "Age at 1 Jan FY"), ("T", "Service (years)"),
 ("U", "Retirement month in FY (0 = none)"), ("V", "Contract end month in FY (0 = none)"), ("W", "Last month on payroll (0 = not active)"),
 ("X", "Increment rate used"), ("Y", "Paid months incl. increment"), ("Z", "Department overtime %"),
 ("AA", "Annual basic (OMR)"), ("AB", "Annual allowances (OMR)"), ("AC", "Overtime (OMR)"), ("AD", "Bonus (OMR)"),
 ("AE", "Social protection (OMR)"), ("AF", "End of service (OMR)"), ("AG", "Tickets, medical & training (OMR)"), ("AH", "FY cost (OMR)"),
 ("AI", "Run-rate cost: 12 months, no increment (OMR)"), ("AJ", "Full-year cost with increment (OMR)"), ("AK", "Monthly gross today (OMR)"),
 ("AL", "Check"), ("AM", "Monthly cost rate before increment (OMR)"), ("AN", "Monthly per-head costs (OMR)"),
]
for i, h in enumerate(DATA_COLS): hdr(x.cell(1, 1 + i), h, "input")
for col, h in CALC_COLS: hdr(x[f"{col}1"], h, "calc")
x.row_dimensions[1].height = 58
deptA, deptB = f"Departments!$A${DEPT0}:$A${DEPTN}", f"Departments!$B${DEPT0}:$B${DEPTN}"
for rr in range(2, LAST + 1):
    n = rr
    F_ = {
     "Q": f'=IF($C{n}="",0,IF(OR(LOWER(TRIM($G{n}))="omani",LOWER(TRIM($G{n}))="oman",LOWER(TRIM($G{n}))="om",LOWER(TRIM($G{n}))="local"),1,0))',
     "R": f'=IF($C{n}="",0,IF(OR(ISNUMBER(SEARCH("resign",$N{n})),ISNUMBER(SEARCH("terminat",$N{n})),ISNUMBER(SEARCH("inactive",$N{n})),ISNUMBER(SEARCH("vacan",$N{n})),ISNUMBER(SEARCH("deceased",$N{n})),ISNUMBER(SEARCH("separated",$N{n}))),0,1))',
     "S": f'=IF(OR($R{n}=0,NOT(ISNUMBER($I{n}))),"",IFERROR(DATEDIF($I{n},DATE(FY,1,1),"y"),""))',
     "T": f'=IF(OR($R{n}=0,NOT(ISNUMBER($J{n}))),"",IFERROR(DATEDIF($J{n},DATE(FY,1,1),"y"),0))',
     "U": f'=IF(OR($R{n}=0,NOT(ISNUMBER($I{n}))),0,IF(YEAR($I{n})+IF($Q{n}=1,RetAgeO,RetAgeE)=FY,MONTH($I{n}),IF(YEAR($I{n})+IF($Q{n}=1,RetAgeO,RetAgeE)<FY,1,0)))',
     "V": f'=IF(OR($R{n}=0,NOT(ISNUMBER($K{n}))),0,IF(YEAR($K{n})=FY,MONTH($K{n}),IF(YEAR($K{n})<FY,1,0)))',
     "W": f'=IF($R{n}=0,0,MIN(12,IF(AND(ISNUMBER($O{n}),$O{n}>=1,$O{n}<=12),INT($O{n}),12),IF(AND(RetireExit="Yes",$U{n}>0),$U{n},12),IF(AND(ContractExit="Yes",$V{n}>0),$V{n},12)))',
     "X": f'=IF(ISNUMBER($P{n}),$P{n},IncPct)',
     "Y": f'=IF($W{n}=0,0,MAX(0,MIN($W{n},IncMonth-1))+MAX(0,$W{n}-IncMonth+1)*(1+$X{n}))',
     "Z": f'=IF($C{n}="",0,IFERROR(INDEX({deptB},MATCH($C{n},{deptA},0)),0))',
     "AA": f'=N($L{n})*$Y{n}',
     "AB": f'=N($M{n})*$Y{n}',
     "AC": f'=$AA{n}*$Z{n}',
     "AD": f'=$AA{n}*BonusMonths/12',
     "AE": f'=($AA{n}+SI_Gross*$AB{n})*IF($Q{n}=1,SI_O,SI_E)',
     "AF": f'=IF($Q{n}=1,0,$AA{n}*EOSB)',
     "AG": f'=$AN{n}*$W{n}',
     "AH": f'=SUM($AA{n}:$AG{n})',
     "AI": f'=12*($AM{n}+$AN{n})',
     "AJ": f'=$AM{n}*((IncMonth-1)+(13-IncMonth)*(1+$X{n}))+12*$AN{n}',
     "AK": f'=IF($R{n}=1,N($L{n})+N($M{n}),0)',
     "AL": f'=IF($C{n}="","",IF(N($L{n})<=0,"Basic salary missing",IF(ISNA(MATCH($C{n},{deptA},0)),"Department not on Departments sheet",IF(AND($I{n}<>"",NOT(ISNUMBER($I{n}))),"Date of birth is text, not a date",IF(AND($K{n}<>"",NOT(ISNUMBER($K{n}))),"Contract end is text, not a date","")))))',
     "AM": f'=IF($R{n}=0,0,N($L{n})*(1+$Z{n}+BonusMonths/12+IF($Q{n}=1,0,EOSB))+N($M{n})+(N($L{n})+SI_Gross*N($M{n}))*IF($Q{n}=1,SI_O,SI_E))',
     "AN": f'=IF($R{n}=0,0,(IF($Q{n}=1,MedO,MedE+Ticket)+Training)/12)',
    }
    for col, fx in F_.items():
        c = x[f"{col}{n}"]; c.value = fx
    for col in ("I", "J", "K"): x[f"{col}{n}"].number_format = "dd/mm/yyyy"
    x[f"L{n}"].number_format = "#,##0.000"; x[f"M{n}"].number_format = "#,##0.000"; x[f"P{n}"].number_format = "0.0%"
    for col in ("AA","AB","AC","AD","AE","AF","AG","AH","AI","AJ","AK","AM","AN"): x[f"{col}{n}"].number_format = HIDE0
    for col in ("Q","R","U","V","W"): x[f"{col}{n}"].number_format = "0;-0;;"
    x[f"X{n}"].number_format = "0.0%"; x[f"Z{n}"].number_format = '0.0%;-0.0%;;'; x[f"Y{n}"].number_format = '0.00;-0.00;;'
widths = [12, 24, 28, 18, 26, 8, 13, 9, 12, 12, 12, 13, 13, 14, 12, 12] + [10]*10 + [13]*11 + [30, 13, 13]
for i, w in enumerate(widths): x.column_dimensions[CL(1 + i)].width = w
x.freeze_panes = "D2"
x.auto_filter.ref = f"A1:AN{LAST}"
x.conditional_formatting.add(f"AL2:AL{LAST}", FormulaRule(formula=['LEN($AL2)>0'], fill=BADF))
x.conditional_formatting.add(f"Q2:AN{LAST}", FormulaRule(formula=['$C2<>""'], fill=CALC))
dv_m = DataValidation(type="whole", operator="between", formula1="1", formula2="12", allow_blank=True, showErrorMessage=True, error="Enter a month number from 1 to 12, or leave blank.")
x.add_data_validation(dv_m); dv_m.add(f"O2:O{LAST}")
x["Q1"].comment = Comment("Columns Q to AN are calculated from columns A to P and the Settings and Departments sheets. Do not type here.", "Workbook")

# ------------------------------------------------------------------ Plan
p = wb.create_sheet("Plan")
title(p, "Plan: new hires and other exits", "Positions not yet in the Data sheet (hires) and leavers who are not named there (attrition). Month = start month for a hire, last month on payroll for an exit.", 20)
PH = [("Type (Hire / Exit)", "input"), ("Department", "input"), ("Position", "input"), ("Nationality", "input"), ("Number", "input"), ("Month (1-12)", "input"),
      ("Basic (OMR / month)", "input"), ("Allowances (OMR / month)", "input"), ("Reason / request reference", "input"),
      ("Omani (1/0)", ""), ("Sign", ""), ("First month affected", ""), ("Last month affected", ""), ("Paid months per head incl. increment", ""),
      ("Department overtime %", ""), ("Months affected", ""), ("Cost per head (OMR)", ""), ("Recruitment (OMR)", ""), ("Total effect (OMR)", ""), ("Heads effect", ""),
      ("Monthly cost rate per head (OMR)", ""), ("Monthly per-head costs (OMR)", "")]
for i, (h, k) in enumerate(PH): hdr(p.cell(4, 1 + i), h, k)
p.row_dimensions[4].height = 58
dv_t = DataValidation(type="list", formula1='"Hire,Exit"', allow_blank=True); p.add_data_validation(dv_t); dv_t.add(f"A{PLAN0}:A{PLANN}")
dv_n = DataValidation(type="list", formula1='"Omani,Expatriate"', allow_blank=True); p.add_data_validation(dv_n); dv_n.add(f"D{PLAN0}:D{PLANN}")
dv_d = DataValidation(type="list", formula1=f"=Departments!$A${DEPT0}:$A${DEPTN}", allow_blank=True); p.add_data_validation(dv_d); dv_d.add(f"B{PLAN0}:B{PLANN}")
dv_pm = DataValidation(type="whole", operator="between", formula1="1", formula2="12", allow_blank=True); p.add_data_validation(dv_pm); dv_pm.add(f"F{PLAN0}:F{PLANN}")
for rr in range(PLAN0, PLANN + 1):
    n = rr
    for c in range(1, 10):
        cell = p.cell(rr, c); cell.font = font(size=10, color=BLUE); cell.border = BOX
    p.cell(rr, 7).number_format = "#,##0"; p.cell(rr, 8).number_format = "#,##0"
    F_ = {
     "J": f'=IF($D{n}="Omani",1,0)',
     "K": f'=IF($A{n}="Hire",1,IF($A{n}="Exit",-1,0))',
     "L": f'=IF($K{n}=1,MAX(1,N($F{n})),IF($K{n}=-1,N($F{n})+1,13))',
     "M": f'=12',
     "N": f'=MAX(0,MIN($M{n},IncMonth-1)-$L{n}+1)+MAX(0,$M{n}-MAX($L{n},IncMonth)+1)*(1+IncPct)',
     "O": f'=IF($B{n}="",0,IFERROR(INDEX({deptB},MATCH($B{n},{deptA},0)),0))',
     "P": f'=MAX(0,$M{n}-$L{n}+1)',
     "Q": f'=$U{n}*$N{n}+$V{n}*$P{n}',
     "R": f'=IF($K{n}=1,N($E{n})*IF($J{n}=1,RecO,RecE),0)',
     "S": f'=$K{n}*N($E{n})*$Q{n}+$R{n}',
     "T": f'=$K{n}*N($E{n})',
     "U": f'=N($G{n})*(1+$O{n}+BonusMonths/12+IF($J{n}=1,0,EOSB))+N($H{n})+(N($G{n})+SI_Gross*N($H{n}))*IF($J{n}=1,SI_O,SI_E)',
     "V": f'=(IF($J{n}=1,MedO,MedE+Ticket)+Training)/12',
    }
    for col, fx in F_.items():
        c = p[f"{col}{n}"]; c.value = fx; c.font = font(size=10)
        c.number_format = HIDE0 if col in ("Q","R","S","U","V") else ('0.00;-0.00;;' if col == "N" else ('0.0%;-0.0%;;' if col == "O" else "0;-0;;"))
p.conditional_formatting.add(f"J{PLAN0}:V{PLANN}", FormulaRule(formula=[f'$A{PLAN0}<>""'], fill=CALC))
p.cell(3, 1, "Plan checks:").font = font(size=9, italic=True)
p.cell(3, 2, f'=IF(SUMPRODUCT(({P("A")}<>"")*(({P("B")}="")+({P("D")}="")+(N(+{P("E")})=0)))>0,"Some plan lines are missing department, nationality or number","All plan lines complete")')
p.cell(3, 2).font = font(size=9, bold=True)
for i, w in enumerate([14, 28, 26, 13, 9, 10, 12, 12, 30] + [11]*13): p.column_dimensions[CL(1 + i)].width = w
p.freeze_panes = "A5"

# ------------------------------------------------------------------ Budget by Dept
b = wb.create_sheet("Budget by Dept")
title(b, "Budget by department and cost line", "Annual OMR. Staff costs are sums over the Data sheet; plan lines and contingency come from Plan and Settings.", 13)
BH = ["Department", "Basic", "Allowances", "Overtime", "Bonus", "Social protection", "End of service", "Tickets, medical & training", "Plan lines (net of exits)", "of which recruitment", "Contingency", "FY budget", "Check vs Departments"]
for i, h in enumerate(BH): hdr(b.cell(4, 1 + i), h)
b.row_dimensions[4].height = 44
for rr in range(DEPT0, DEPTN + 1):
    A = f"$A{rr}"
    b.cell(rr, 1, f'=IF(Departments!A{rr}="","",Departments!A{rr})').font = font(size=10, color=GREEN)
    for j, col in enumerate(["AA", "AB", "AC", "AD", "AE", "AF", "AG"]):
        c = b.cell(rr, 2 + j, f'=IF({A}="","",SUMIFS({D(col)},{D("C")},{A}))'); c.number_format = OMR; c.font = font(size=10)
    b.cell(rr, 9, f'=IF({A}="","",SUMIFS({P("S")},{P("B")},{A}))').number_format = OMR
    b.cell(rr, 10, f'=IF({A}="","",SUMIFS({P("R")},{P("B")},{A}))').number_format = OMR
    b.cell(rr, 11, f'=IF({A}="","",SUM(B{rr}:I{rr})*Contingency)').number_format = OMR
    b.cell(rr, 12, f'=IF({A}="","",SUM(B{rr}:I{rr})+K{rr})').number_format = OMR
    b.cell(rr, 13, f'=IF({A}="","",IF(ROUND(L{rr}-Departments!S{rr},0)=0,"OK","Differs"))')
    for c in range(9, 14): b.cell(rr, c).font = font(size=10)
    b.cell(rr, 12).font = font(size=10, bold=True)
b.cell(TR, 1, "Total").font = font(bold=True, size=10)
for c in range(2, 13):
    L = CL(c); cell = b.cell(TR, c, f"=SUM({L}{DEPT0}:{L}{DEPTN})"); cell.number_format = OMR; cell.font = font(bold=True, size=10); cell.border = BOX
b.conditional_formatting.add(f"M{DEPT0}:M{DEPTN}", CellIsRule(operator="equal", formula=['"Differs"'], fill=BADF))
b.column_dimensions["A"].width = 32
for c in range(2, 14): b.column_dimensions[CL(c)].width = 14
b.freeze_panes = "B5"

# ------------------------------------------------------------------ Monthly
m = wb.create_sheet("Monthly")
title(m, "Monthly phasing", "Cost and headcount by month, for cash flow and payroll forecasting. Includes contingency.", 14)
hdr(m.cell(4, 1), "")
for k in range(1, 13):
    hdr(m.cell(4, 1 + k), ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"][k-1])
    m.cell(3, 1 + k, k).font = font(size=8, color="A6A6A6")
hdr(m.cell(4, 14), "Total / check")
labels = ["Staff cost (OMR)", "Plan lines cost (OMR)", "Recruitment (OMR)", "Contingency (OMR)", "Total cost (OMR)", "Headcount (staff)", "Headcount change from plan", "Headcount (total)"]
for i, lab in enumerate(labels): m.cell(5 + i, 1, lab).font = font(size=10, bold=(lab.startswith("Total") or lab == "Headcount (total)"))
for k in range(1, 13):
    col = CL(1 + k); mm = f"{col}$3"
    m[f"{col}5"] = f'=SUMPRODUCT(({D("W")}>={mm})*({D("AM")}*(1+{D("X")}*({mm}>=IncMonth))+{D("AN")}))'
    m[f"{col}6"] = f'=SUMPRODUCT({P("T")}*({mm}>={P("L")})*({mm}<={P("M")})*({P("U")}*(1+IncPct*({mm}>=IncMonth))+{P("V")}))'
    m[f"{col}7"] = f'=SUMPRODUCT(({P("A")}="Hire")*({P("L")}={mm})*{P("R")})'
    m[f"{col}8"] = f'=SUM({col}5:{col}7)*Contingency'
    m[f"{col}9"] = f'=SUM({col}5:{col}8)'
    m[f"{col}10"] = f'=SUMPRODUCT(({D("W")}>={mm})*{D("R")})'
    m[f"{col}11"] = f'=SUMPRODUCT({P("T")}*({mm}>={P("L")})*({mm}<={P("M")}))'
    m[f"{col}12"] = f'={col}10+{col}11'
    for r_ in range(5, 13):
        m[f"{col}{r_}"].number_format = OMR if r_ <= 9 else INT; m[f"{col}{r_}"].font = font(size=10, bold=(r_ in (9, 12)))
for r_ in range(5, 10):
    m[f"N{r_}"] = f"=SUM(B{r_}:M{r_})"; m[f"N{r_}"].number_format = OMR; m[f"N{r_}"].font = font(size=10, bold=(r_ == 9))
m["A14"] = "Check: monthly total equals FY budget"; m["A14"].font = font(size=10, italic=True)
m["N14"] = f'=IF(ROUND(N9-Departments!S{TR},0)=0,"OK","Differs by "&TEXT(N9-Departments!S{TR},"#,##0"))'
m.column_dimensions["A"].width = 32
for k in range(2, 15): m.column_dimensions[CL(k)].width = 12
m.column_dimensions["N"].width = 16
lc = LineChart(); lc.title = "Total cost by month (OMR)"; lc.height = 7; lc.width = 22; lc.legend = None
lc.add_data(Reference(m, min_col=2, max_col=13, min_row=9), from_rows=True, titles_from_data=False)
lc.set_categories(Reference(m, min_col=2, max_col=13, min_row=4)); lc.y_axis.numFmt = "#,##0"; lc.y_axis.title = "OMR"
m.add_chart(lc, "A17")

# ------------------------------------------------------------------ Overview (first visible sheet after Guide)
o = wb.create_sheet("Overview", 1)
title(o, "Overview", "Calculated from the Data, Plan, Departments and Settings sheets. Nothing on this sheet is typed in.", 9)
def kv(r_, label, fx, fmt=OMR, bold=False, note=None):
    o.cell(r_, 1, label).font = font(size=10, bold=bold)
    c = o.cell(r_, 2, fx); c.number_format = fmt; c.font = font(size=11 if bold else 10, bold=bold); c.border = BOX
    if note: o.cell(r_, 3, note).font = font(size=9, italic=True, color="595959")
o["A4"] = "Data source"; o["A4"].font = font(bold=True, size=12, color="1F3864")
kv(5, "Rows in the Data sheet", f'=COUNTA({D("C")})', INT)
kv(6, "Active employees", f'=SUM({D("R")})', INT)
kv(7, "Rows with data issues (fix on Data, column AL)", f'=COUNTIF({D("AL")},"Basic*")+COUNTIF({D("AL")},"Department*")+COUNTIF({D("AL")},"Date*")+COUNTIF({D("AL")},"Contract*")', INT)
kv(8, "Assumptions not yet confirmed (Settings)", "=Unconfirmed", INT)
o["C5"] = '=IF(B5=0,"No data yet: paste your employee export into the Data sheet. All figures below stay at zero until then.","")'
o["C5"].font = font(size=10, bold=True, color="C00000")
o["A10"] = "Key figures"; o["A10"].font = font(bold=True, size=12, color="1F3864")
kv(11, "FY budget (OMR)", f"=Departments!S{TR}", OMR, True)
kv(12, "Prior-year budget (OMR)", f"=Departments!C{TR}", OMR)
kv(13, "Change vs prior year", f'=IF(B12=0,"",B11/B12-1)', PCT)
kv(14, "Headcount today", f"=Departments!E{TR}", INT)
kv(15, "Closing headcount", f"=Departments!O{TR}", INT)
kv(16, "Omanisation today", f'=IF(B6=0,"",SUMPRODUCT({D("Q")},{D("R")})/B6)', PCT)
kv(17, "Omanisation at year end", f"=Departments!V{TR}", PCT)
kv(18, "Omanisation target", "=OmanTarget", PCT)
kv(19, "Average FTE in the year", f'=SUM({D("W")})/12+SUMPRODUCT({P("T")},{P("P")})/12', '#,##0.0')
kv(20, "Average cost per FTE (OMR)", '=IF(N(B19)=0,"",B11/B19)', OMR)
o["C17"] = '=IF(OR(B17="",B18=""),"",IF(B17>=B18,"Meets target","Below target"))'
o["A22"] = "Budget bridge"; o["A22"].font = font(bold=True, size=12, color="1F3864")
for i, h in enumerate(["Step", "OMR", "Explanation", "", "Base", "Increase", "Decrease"]): hdr(o.cell(23, 1 + i), h)
steps = [
 ("Payroll run-rate today", f'=SUM({D("AI")})', "Active staff at today's pay for 12 months, with all on-costs, no increment."),
 ("Salary increment", f'=SUM({D("AJ")})-SUM({D("AI")})', "Increment from its effective month (personal rates where given)."),
 ("Named leavers", f'=SUM({D("AH")})-SUM({D("AJ")})', "Planned exits, and retirements or contract ends when Settings says to treat them as exits."),
 ("Planned hires", f'=SUMIFS({P("S")},{P("A")},"Hire")', "New positions on the Plan sheet, including recruitment."),
 ("Other exits (attrition)", f'=SUMIFS({P("S")},{P("A")},"Exit")', "Exits on the Plan sheet."),
 ("Contingency", f"=Departments!R{TR}", "Contingency % on each department's total."),
 ("FY budget", "=SUM(B24:B29)", "Sum of the steps above."),
]
for i, (lab, fx, ex) in enumerate(steps):
    r_ = 24 + i
    o.cell(r_, 1, lab).font = font(size=10, bold=(i in (0, 6)))
    c = o.cell(r_, 2, fx); c.number_format = OMR; c.font = font(size=10, bold=(i in (0, 6))); c.border = BOX
    o.cell(r_, 3, ex).font = font(size=9, color="595959")
    # waterfall helper columns E:G (base / increase / decrease)
    if i == 0 or i == 6:
        o.cell(r_, 5, 0); o.cell(r_, 6, f"=B{r_}"); o.cell(r_, 7, 0)
    else:
        prev = f"SUM($B$24:B{r_-1})"
        o.cell(r_, 5, f"=MIN({prev},{prev}+B{r_})"); o.cell(r_, 6, f"=MAX(B{r_},0)"); o.cell(r_, 7, f"=MAX(-B{r_},0)")
    for cc in (5, 6, 7): o.cell(r_, cc).number_format = OMR; o.cell(r_, cc).font = font(size=8, color="A6A6A6")
o["A31"] = "Check: bridge equals budget"; o["A31"].font = font(size=10, italic=True)
o["B31"] = f'=IF(ROUND(B30-B11,0)=0,"OK","Differs by "&TEXT(B30-B11,"#,##0")&": a department in Data is missing from Departments")'
o["A33"] = "Reconciliation with the Data sheet"; o["A33"].font = font(bold=True, size=12, color="1F3864")
for i, h in enumerate(["Measure", "Data sheet total", "By department", "Difference", "Status"]): hdr(o.cell(34, 1 + i), h)
rec = [("Active headcount", f'=SUM({D("R")})', f"=Departments!E{TR}", INT),
       ("Monthly gross pay today (OMR)", f'=SUM({D("AK")})', f"=Departments!I{TR}", OMR),
       ("Staff FY cost (OMR)", f'=SUM({D("AH")})', f"=Departments!P{TR}", OMR),
       ("FY budget vs monthly phasing (OMR)", "=Monthly!N9", f"=Departments!S{TR}", OMR)]
for i, (lab, a, bb, fmt) in enumerate(rec):
    r_ = 35 + i
    o.cell(r_, 1, lab).font = font(size=10)
    for cc, fx in ((2, a), (3, bb), (4, f"=B{r_}-C{r_}")):
        c = o.cell(r_, cc, fx); c.number_format = fmt; c.font = font(size=10); c.border = BOX
    o.cell(r_, 5, f'=IF(ROUND(D{r_},0)=0,"✓ Matches","✗ Check")').font = font(size=10, bold=True)
o.conditional_formatting.add("E35:E38", FormulaRule(formula=['LEFT(E35,1)="✗"'], fill=BADF))
o.conditional_formatting.add("E35:E38", FormulaRule(formula=['LEFT(E35,1)="✓"'], fill=OKF))
o.conditional_formatting.add("B7:B8", CellIsRule(operator="greaterThan", formula=["0"], fill=WARNF))
o.conditional_formatting.add("C17", CellIsRule(operator="equal", formula=['"Below target"'], fill=WARNF))
o.conditional_formatting.add("B31", FormulaRule(formula=['LEFT(B31,2)<>"OK"'], fill=BADF))
o.column_dimensions["A"].width = 44; o.column_dimensions["B"].width = 18; o.column_dimensions["C"].width = 64
for c in "DEFG": o.column_dimensions[c].width = 13
# waterfall chart
wf = BarChart(); wf.type = "col"; wf.grouping = "stacked"; wf.overlap = 100; wf.title = "Budget bridge (OMR)"; wf.height = 8; wf.width = 22
wf.add_data(Reference(o, min_col=5, max_col=7, min_row=23, max_row=30), titles_from_data=True)
wf.set_categories(Reference(o, min_col=1, min_row=24, max_row=30))
wf.series[0].graphicalProperties.noFill = True; wf.series[0].graphicalProperties.line.noFill = True
wf.series[1].graphicalProperties.solidFill = "2A78D6"; wf.series[2].graphicalProperties.solidFill = "1F7A55"
wf.y_axis.numFmt = "#,##0"; wf.legend = None
o.add_chart(wf, "I4")
# budget by department chart (Departments column S)
bc = BarChart(); bc.type = "bar"; bc.title = "FY budget by department (OMR)"; bc.height = 10; bc.width = 22; bc.legend = None
bc.add_data(Reference(d, min_col=19, min_row=DEPT0, max_row=min(DEPTN, DEPT0 + 19)), titles_from_data=False)
bc.set_categories(Reference(d, min_col=1, min_row=DEPT0, max_row=min(DEPTN, DEPT0 + 19)))
bc.series[0].graphicalProperties.solidFill = "2A78D6"; bc.x_axis.scaling.orientation = "maxMin"; bc.y_axis.numFmt = "#,##0"
o.add_chart(bc, "I22")
o.sheet_properties.tabColor = "1F3864"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
